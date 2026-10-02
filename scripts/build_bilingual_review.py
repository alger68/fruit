"""Build docs/Global_AR_中英文對照審閱表.xlsx: Chinese and English wording side by side for review."""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
from translations_en import DESCRIPTION_EN, EN, PROFILE_EN, SECTIONS_EN  # noqa: E402

zh = json.load(open("shared/survey-questions.json", encoding="utf-8"))["questions"]
HDR = PatternFill("solid", fgColor="1F3A5F")
QFILL = PatternFill("solid", fgColor="DDE7F3")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BBBBBB")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def put(ws, r, vals, fill=None, bold=False):
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v)
        x.alignment, x.border = WRAP, B
        if fill:
            x.fill = fill
        if bold:
            x.font = Font(bold=True)


def head(ws, heads):
    for c, h in enumerate(heads, 1):
        x = ws.cell(1, c, h)
        x.fill, x.font, x.border = HDR, Font(bold=True, color="FFFFFF"), B
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


wb = Workbook()
ws0 = wb.active
ws0.title = "說明 Notes"
notes = [
    "Global AR 問卷:中英文對照審閱表 / Chinese-English review sheet",
    "",
    "用途:讓審閱者並排比對中文與英文。表單上每題與每個選項會顯示為「中文 / English」。",
    "Purpose: compare the Chinese and English wording side by side. The form shows each question and option as \"中文 / English\".",
    "",
    "請在黃色欄位填寫。「審閱結果」選 OK 或需修改,「建議英文用詞」寫你希望的說法。",
    "Please use the yellow columns: pick OK or Needs change in the result column and write your preferred wording in the suggestion column.",
    "",
    "選項若是純英文(例如 Cash、ACH),表單上不會再加中文。",
    "Options that are already English only (for example Cash, ACH) are shown without Chinese.",
    "",
    "表單建立後,請不要手動改選項文字,統計表是依整串文字比對。",
    "After the form is created, do not edit option text by hand. The statistics sheet matches the full text.",
]
for i, t in enumerate(notes, 1):
    c = ws0.cell(i, 1, t)
    c.alignment = WRAP
ws0["A1"].font = Font(bold=True, size=14)
ws0.column_dimensions["A"].width = 120

ws = wb.create_sheet("題目對照 Questions")
head(ws, ["題號 ID", "區塊 Section", "類型 Type", "中文 Chinese", "English", "題型 Kind", "必填 Req.", "審閱結果 Result", "建議英文用詞 Suggested wording"])
r = 2
put(ws, r, ["—", "Form", "說明 Description", "為建立一致的全球應收帳款(AR)作業標準,請協助填寫本問卷。全部為選擇題,約需 20 至 30 分鐘。標 * 為必填。若選項無法描述貴單位的做法,請選「其他」並簡短說明。",
            DESCRIPTION_EN, "", "", "", ""], fill=QFILL)
r += 1
for pid, zhs in [("P1", "夥伴代碼(請填寫總部提供的代碼)"), ("P2", "公司名稱"), ("P3", "國家 / 地區"), ("P4", "聯絡 Email(選填,供補問使用)")]:
    put(ws, r, [pid, "Page 1", "題目 Question", zhs, PROFILE_EN[pid], "文字 Text", "是 Yes" if pid != "P4" else "否 No", "", ""], fill=QFILL, bold=True)
    r += 1
cur_sec = None
for q in zh:
    if q["id"] != "R0" and q["section"] != cur_sec:
        cur_sec = q["section"]
        put(ws, r, ["", cur_sec, "區塊 Section", cur_sec, SECTIONS_EN[cur_sec[0]], "", "", "", ""], fill=PatternFill("solid", fgColor="C9D9EC"), bold=True)
        r += 1
    en_text, en_opts = EN[q["id"]]
    req = "是 Yes" if (q["required"] or q["id"] == "R0") else "否 No"
    kind = "多選 Multiple" if q["type"] == "multi" else "單選 Single"
    put(ws, r, [q["id"], "Page 1" if q["id"] == "R0" else q["section"].split(" ")[0], "題目 Question", q["text"], en_text, kind, req, "", ""], fill=QFILL, bold=True)
    r += 1
    for z, e in zip(q["options"], en_opts):
        put(ws, r, [q["id"], "", "選項 Option", z, e, "", "", "", ""])
        r += 1
    if q["allowOther"]:
        put(ws, r, [q["id"], "", "選項 Option", "其他(請說明)", "Other (please specify)", "", "", "", ""])
        r += 1
for rr in range(2, r):
    for c in (8, 9):
        ws.cell(rr, c).fill = INPUT
dv = DataValidation(type="list", formula1='"OK,需修改 Needs change"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"H2:H{r - 1}")
for c, w in enumerate([8, 18, 14, 60, 66, 14, 9, 16, 44], 1):
    ws.column_dimensions[chr(64 + c)].width = w
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:I{r - 1}"

wg = wb.create_sheet("術語表 Glossary")
head(wg, ["中文 Chinese", "English", "說明 Note", "審閱結果 Result", "建議用詞 Suggested wording"])
terms = [
    ("銷帳", "Cash application", "把收到的款項對應到應收帳款並沖銷"),
    ("沖帳", "Offset / reversal", "已銷帳款項跨 B/L、帳單或客戶調整"),
    ("放貨", "Cargo release", "進口貨物放行"),
    ("放單", "Document release (B/L)", "出口提單放行"),
    ("Hold", "Hold", "貨物或單據的扣留狀態"),
    ("授信客戶", "Credit customer", "給予付款期限的客戶"),
    ("第三地付款", "Third-location payment", "由第三地或第三方付款的款項"),
    ("到付 / 運費到付", "Freight collect", "運費由收貨人支付"),
    ("匯兌盈損", "FX gain / loss", "匯率差異造成的盈虧"),
    ("航次匯率", "Voyage rate", "以航次為基準的匯率"),
    ("短付 / 溢付", "Short payment / overpayment", "付款少於或多於應付金額"),
    ("折讓", "Discount", ""),
    ("減免", "Waiver (write-off)", "小額差異免收"),
    ("預收款", "Prepayment", "尚未對應到帳單的先收款"),
    ("未辨識款", "Unidentified receipt", "收到但無法對應客戶或帳單的款項"),
    ("訂艙 / 開船 / 到港", "Booking / Sailing / Arrival", ""),
    ("提單", "Bill of lading (B/L)", ""),
    ("發票 / 帳單", "Invoice / billing statement", "航運業常以 Arrival Notice 或提單影本代替發票"),
    ("報收", "LFC reporting", "向總部報送應收或收款資料"),
    ("在地", "Local office", "代理行所在地的作業單位"),
    ("總部", "HQ", ""),
    ("代理行", "Agent", ""),
    ("子公司 / 合資公司", "Subsidiary / joint venture", ""),
    ("核准層級", "Approval level", ""),
]
for i, t in enumerate(terms, 2):
    put(wg, i, list(t) + ["", ""])
    for c in (4, 5):
        wg.cell(i, c).fill = INPUT
dv2 = DataValidation(type="list", formula1='"OK,需修改 Needs change"', allow_blank=True)
wg.add_data_validation(dv2)
dv2.add(f"D2:D{len(terms) + 1}")
for c, w in enumerate([24, 32, 50, 16, 36], 1):
    wg.column_dimensions[chr(64 + c)].width = w
wg.freeze_panes = "A2"

out = "docs/Global_AR_中英文對照審閱表.xlsx"
wb.save(out)
print(out, r - 1, "rows")
